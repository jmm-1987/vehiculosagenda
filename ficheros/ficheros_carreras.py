from flask import request, send_file
import openpyxl
import io
import pandas as pd
from openpyxl.styles import Font
import zipfile

def register_func_subir_fichero_carreras(app):
    @app.route("/procesar_carreras", methods=['POST'])
    def procesar_carreras():
        ref_file = request.files.get('fichero_carreras')
        if not ref_file:
            return "No se ha subido el fichero de referencia", 400
        comp_files = request.files.getlist('ficheros_comparacion')
        if not comp_files or len(comp_files) == 0:
            return "No se han subido ficheros de comparación", 400
        if len(comp_files) > 6:
            return "Solo se permiten hasta 6 ficheros de comparación", 400

        # Leer referencia
        ref_df = pd.read_excel(ref_file)
        ref_agencias_set = set(ref_df['Ref.Agencia'].astype(str).str.strip())
        # Leer todos los albaranes e importes de los ficheros de comparación
        comp_albaranes = set()
        comp_importes = dict()
        comp_extra_rows = []
        for file in comp_files:
            df = pd.read_excel(file)
            for _, row in df.iterrows():
                albaran_comp = str(row.get('Albarán', '')).strip()
                comp_albaranes.add(albaran_comp)
                importe = row.get('Importe', None)
                if pd.notnull(importe):
                    comp_importes[albaran_comp] = importe
                # Guardar filas que no están en el de referencia
                if albaran_comp and albaran_comp not in ref_agencias_set:
                    comp_extra_rows.append(row)

        # Insertar columna COBRADO tras Ref.Agencia y columna Importe Cobrado tras T.Venta
        cols = list(ref_df.columns)
        idx_refagencia = cols.index('Ref.Agencia') + 1
        idx_tventa = cols.index('T.Venta') + 1
        ref_df.insert(idx_refagencia, 'COBRADO', ref_df['Ref.Agencia'].astype(str).apply(lambda x: 'SÍ' if x.strip() in comp_albaranes else 'NO'))
        ref_df.insert(idx_tventa + 1, 'Importe Cobrado', ref_df['Ref.Agencia'].astype(str).apply(lambda x: comp_importes.get(x.strip(), '')))

        # Guardar a Excel y poner en negrita toda la columna de las columnas nuevas
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            ref_df.to_excel(writer, index=False)
        output.seek(0)
        wb = openpyxl.load_workbook(output)
        ws = wb.active
        col_indices = {cell.value: cell.column for cell in ws[1] if cell.value in ['COBRADO', 'Importe Cobrado']}
        for col in col_indices.values():
            for row in ws.iter_rows(min_row=1, min_col=col, max_col=col, max_row=ws.max_row):
                for cell in row:
                    cell.font = Font(bold=True)
        final_output = io.BytesIO()
        wb.save(final_output)
        final_output.seek(0)

        # Crear el segundo Excel con las filas extra
        extra_output = io.BytesIO()
        if comp_extra_rows:
            extra_df = pd.DataFrame(comp_extra_rows)
            with pd.ExcelWriter(extra_output, engine='openpyxl') as writer:
                extra_df.to_excel(writer, index=False)
            extra_output.seek(0)
        else:
            # Si no hay filas extra, crear un Excel vacío con mensaje
            with pd.ExcelWriter(extra_output, engine='openpyxl') as writer:
                pd.DataFrame({'Mensaje': ['No hay expediciones extra encontradas.']}).to_excel(writer, index=False)
            extra_output.seek(0)

        # Crear un ZIP con ambos archivos
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w') as zipf:
            zipf.writestr('Carreras_comparadas.xlsx', final_output.getvalue())
            zipf.writestr('Expediciones_extra.xlsx', extra_output.getvalue())
        zip_buffer.seek(0)
        return send_file(zip_buffer, download_name="Carreras_resultados.zip", as_attachment=True) 
from flask import Flask, request, send_file
import fitz
import re
import io
import openpyxl
from openpyxl.styles import numbers
from collections import defaultdict
import zipfile

app = Flask(__name__)

def register_func_subir_fichero_pallex(app):
    @app.route("/subir_fichero_pallex", methods=['POST'])
    def subir_fichero_pallex():
        files = request.files.getlist('ficheros')
        if not files or len(files) == 0:
            return "No se ha subido ningún archivo", 400

        agrupado = defaultdict(lambda: defaultdict(float))
        tipos_set = set()

        for file in files:
            wb = openpyxl.load_workbook(file, data_only=True)
            ws = wb.active
            for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
                tipo = row[1].value
                importe = row[2].value
                num_envio = row[4].value  # Columna E
                if num_envio is None or tipo is None or importe is None:
                    continue
                tipos_set.add(str(tipo))
                agrupado[str(num_envio)][str(tipo)] += float(importe)

        tipos_ordenados = sorted(tipos_set)
        columnas = ["Nº envío"] + tipos_ordenados + ["Total"]

        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = "Resumen Pallex"
        ws_out.append(columnas)

        for num_envio, tipos_dict in agrupado.items():
            fila = [num_envio]
            total = 0.0
            for tipo in tipos_ordenados:
                valor = tipos_dict.get(tipo, 0.0)
                fila.append(valor)
                total += valor
            fila.append(total)
            ws_out.append(fila)

        output = io.BytesIO()
        wb_out.save(output)
        output.seek(0)
        # Llamar a la función para exportar el segundo Excel
        conceptos_output = exportar_excel_conceptos(agrupado, tipos_ordenados)
        # Crear un zip con ambos archivos
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w') as zipf:
            zipf.writestr('Resumen_Pallex.xlsx', output.getvalue())
            zipf.writestr('Conceptos_Pallex.xlsx', conceptos_output.getvalue())
        zip_buffer.seek(0)
        return send_file(zip_buffer, download_name="Pallex_resultados.zip", as_attachment=True)

def exportar_excel_conceptos(agrupado, tipos_ordenados):
    import openpyxl
    import io
    clave_map = {
        'Coste de Hub': 'H',
        'Entrega': 'E',
        'Fondo de contingencia en origen': 'C',
        'Recogida': 'R',
        'Reentrega': 'R2',
        'Seguro MET Origen': 'S',
    }
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conceptos Pallex"
    ws.append(["", "", "", "", "Cantidad", "Clave", "", "", "Nº envío"])
    for num_envio, tipos_dict in agrupado.items():
        for tipo, cantidad in tipos_dict.items():
            clave = clave_map.get(tipo, tipo)
            ws.append(["", "", "", "", cantidad, clave, "", "", num_envio])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
